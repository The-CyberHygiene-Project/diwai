"""Readers, chunking, dedupe, ingest/list/delete. The only code path that writes chunks.

CLI:  ingest.py <file|folder> [...] [--force]
      ingest.py --list
      ingest.py --delete NAME
One writer: don't run this while the browser manager is adding files.
"""
import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime

import config
import rag_utils

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".pptx", ".htm", ".html", ".dita"}
SIDECAR_SUFFIX = ".source.json"
NEAR_DUP_COSINE = 0.99
NO_TEXT_PDF = "No text found (needs OCR)"


@dataclass
class IngestResult:
    status: str  # ingested | skipped_duplicate | empty | error
    source: str
    chunks: int = 0
    duplicate_of: str = ""
    detail: str = ""
    stage: str = ""  # "embed" when the embedder, not the file, failed


class EmbedFailed(Exception):
    """The embedding server failed; the file itself may be fine."""


class UnsupportedType(Exception):
    pass


# ---- identity ---------------------------------------------------------------

def doc_id_for(path):
    real = os.path.realpath(os.path.abspath(path))
    return hashlib.sha1(real.encode("utf-8")).hexdigest()[:16]


def normalized_sha256(text):
    return hashlib.sha256(" ".join(text.split()).encode("utf-8")).hexdigest()


# ---- readers ----------------------------------------------------------------

def _pdf_text_pypdf(path):
    from pypdf import PdfReader
    return "\n\n".join((page.extract_text() or "") for page in PdfReader(path).pages)


def _pdf_text_gs(path):
    out = subprocess.run(
        ["gs", "-q", "-dSAFER", "-dNOPAUSE", "-dBATCH", "-sDEVICE=txtwrite",
         "-sOutputFile=-", os.path.abspath(path)],
        capture_output=True, timeout=600,
    )
    return out.stdout.decode("utf-8", errors="replace")


def _read_pdf(path):
    try:
        text = _pdf_text_pypdf(path)
    except Exception:
        text = ""
    if not text.strip():
        text = _pdf_text_gs(path)
    return text


def _read_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(path)
    parts = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells, seen = [], set()
                for cell in row.cells:  # merged cells repeat; keep one
                    if id(cell._tc) not in seen:
                        seen.add(id(cell._tc))
                        cells.append(cell.text.strip())
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pptx(path):
    from pptx import Presentation

    parts = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        parts.append(f"[Slide {n}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


_HTML_DROP = ("//script", "//style", "//noscript", "//template")
_HTML_CHROME = ("nav", "aside", "footer", "*[@role='navigation']",
                "*[@role='complementary']", "*[@role='contentinfo']")
_HTML_BLOCKS = {"p", "div", "section", "article", "main", "h1", "h2", "h3", "h4", "h5", "h6",
                "pre", "table", "tr", "ul", "ol", "dl", "dt", "dd", "blockquote", "figure"}


def _read_html(path):
    """Vendor pages: keep the main content, drop site navigation, headers,
    footers and side panels, and keep paragraph breaks for the chunker."""
    import lxml.html

    parser = lxml.html.HTMLParser(no_network=True, remove_comments=True)
    tree = lxml.html.parse(path, parser).getroot()
    if tree is None:
        return ""
    for xp in _HTML_DROP:
        for el in tree.xpath(xp):
            el.drop_tree()
    mains = tree.xpath("//main | //*[@role='main']")
    articles = [] if mains else tree.xpath("//article[not(ancestor::article)]")
    if mains:
        root = mains[0]
    elif articles:  # a page of several articles: keep every one (finding 8)
        root = lxml.html.Element("div")
        root.extend(articles)
    else:
        root = tree.find("body") if tree.find("body") is not None else tree
    mains = mains or articles
    chrome = list(_HTML_CHROME) + ([] if mains else ["header", "*[@role='banner']"])
    for tag in chrome:
        for el in root.xpath(f".//{tag}"):
            el.drop_tree()
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        if el.tag in _HTML_BLOCKS:
            el.tail = "\n\n" + (el.tail or "")
        elif el.tag in ("li", "br"):
            el.tail = "\n" + (el.tail or "")
    text = root.text_content()
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _read_dita(path):
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True,
                             load_dtd=False, remove_comments=True)
    root = etree.parse(path, parser).getroot()
    # itertext keeps <title> text, which carries the identifier.
    return "\n".join(t.strip() for t in root.itertext() if t.strip())


def read_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext in (".htm", ".html"):
        return _read_html(path)
    if ext == ".dita":
        return _read_dita(path)
    raise UnsupportedType(f"Unsupported file type: {ext or '(none)'}")


# ---- provenance -------------------------------------------------------------

def provenance_for(path):
    """Fields from <file>.source.json (URL, fetch date, sha256), verified
    against the file's bytes. {} when there is no sidecar."""
    sidecar = path + SIDECAR_SUFFIX
    if not os.path.exists(sidecar):
        return {}
    with open(sidecar, encoding="utf-8") as fh:
        side = json.load(fh)
    with open(path, "rb") as fh:
        actual = hashlib.sha256(fh.read()).hexdigest()
    return {
        "provenance_url": str(side.get("url", "")),
        "provenance_fetched": str(side.get("fetched", "")),
        "provenance_sha256": str(side.get("sha256", "")),
        "provenance_verified": side.get("sha256") == actual,
    }


# ---- dedupe -----------------------------------------------------------------

def _exact_duplicate(col, sha, doc_id):
    got = col.get(where={"content_sha256": sha}, include=["metadatas"])
    for meta in got["metadatas"]:
        if meta.get("doc_id") != doc_id:
            return meta.get("source", "?")
    return ""


def _near_duplicate(col, embeddings, doc_id):
    """Duplicate only if first, middle and last chunk all hit the SAME other
    document at cosine >= NEAR_DUP_COSINE."""
    picks = [embeddings[0], embeddings[len(embeddings) // 2], embeddings[-1]]
    res = col.query(query_embeddings=picks, n_results=1,
                    where={"doc_id": {"$ne": doc_id}},
                    include=["metadatas", "distances"])
    hits = set()
    for metas, dists in zip(res["metadatas"], res["distances"]):
        if not metas or (1 - dists[0]) < NEAR_DUP_COSINE:
            return ""
        hits.add((metas[0].get("doc_id"), metas[0].get("source", "?")))
    return hits.pop()[1] if len(hits) == 1 else ""


# ---- ingest -----------------------------------------------------------------

def ingest_file(path, collection, force=False, progress=None):
    source = os.path.basename(path)
    try:
        return _ingest(path, source, collection, force, progress)
    except EmbedFailed as exc:
        cause = exc.__cause__ or exc
        return IngestResult("error", source, stage="embed",
                            detail=f"{type(cause).__name__}: {cause}")
    except Exception as exc:  # reported, never raised: one bad file can't stop a batch
        return IngestResult("error", source, detail=f"{type(exc).__name__}: {exc}")


def _ingest(path, source, col, force, progress):
    text = read_file(path)
    if not text.strip():
        detail = NO_TEXT_PDF if path.lower().endswith(".pdf") else "No text found"
        return IngestResult("empty", source, detail=detail)

    doc_id = doc_id_for(path)
    sha = normalized_sha256(text)
    if not force:
        dup = _exact_duplicate(col, sha, doc_id)
        if dup:
            return IngestResult("skipped_duplicate", source, duplicate_of=dup,
                                detail=f"Same text as {dup}")

    chunks = rag_utils.chunk_text(text)
    # Embed first: an embedding failure must never lose the old copy.
    embeddings = []
    for i in range(0, len(chunks), config.EMBED_BATCH):
        try:
            embeddings.extend(rag_utils.embed_documents(chunks[i:i + config.EMBED_BATCH]))
        except Exception as exc:
            raise EmbedFailed() from exc
        if progress:
            progress(len(embeddings), len(chunks))

    if not force and col.count():
        dup = _near_duplicate(col, embeddings, doc_id)
        if dup:
            return IngestResult("skipped_duplicate", source, duplicate_of=dup,
                                detail=f"Nearly identical to {dup}")

    base = {
        "source": source,
        "file_path": os.path.realpath(os.path.abspath(path)),
        "doc_id": doc_id,
        "content_sha256": sha,
        "added_at": datetime.now().isoformat(timespec="seconds"),
        **provenance_for(path),
    }
    # Write the new chunks over the old ones first (ids are deterministic),
    # then drop only the old tail, so a failure part-way never leaves the
    # document missing (finding 4).
    ids = [f"{doc_id}::chunk::{i}" for i in range(len(chunks))]
    metas = [{**base, "chunk_index": i} for i in range(len(chunks))]
    step = config.UPSERT_BATCH
    for i in range(0, len(chunks), step):
        col.upsert(ids=ids[i:i + step], documents=chunks[i:i + step],
                   embeddings=embeddings[i:i + step], metadatas=metas[i:i + step])
    col.delete(where={"$and": [{"doc_id": doc_id}, {"chunk_index": {"$gte": len(chunks)}}]})
    return IngestResult("ingested", source, chunks=len(chunks))


# ---- list / delete ----------------------------------------------------------

def list_documents(col):
    docs = {}
    offset = 0
    while True:
        page = col.get(include=["metadatas"], limit=config.UPSERT_BATCH, offset=offset)
        if not page["ids"]:
            break
        for meta in page["metadatas"]:
            d = docs.setdefault(meta["doc_id"], {
                "doc_id": meta["doc_id"],
                "source": meta.get("source", "?"),
                "file_path": meta.get("file_path", ""),
                "added_at": meta.get("added_at", ""),
                "chunks": 0,
            })
            d["chunks"] += 1
        offset += len(page["ids"])
    return sorted(docs.values(), key=lambda d: (d["source"].lower(), d["file_path"]))


def delete_document(col, name):
    """Remove every document whose file name (or full path) is NAME.
    Returns the number of documents removed."""
    real = os.path.realpath(os.path.abspath(name))
    targets = {d["doc_id"] for d in list_documents(col)
               if d["source"] == name or d["file_path"] == real}
    for doc_id in targets:
        col.delete(where={"doc_id": doc_id})
    return len(targets)


# ---- CLI --------------------------------------------------------------------

def _files_under(path):
    if os.path.isfile(path):
        yield path
        return
    for root, dirs, files in os.walk(path):
        dirs[:] = sorted(d for d in dirs if not d.startswith("."))
        for f in sorted(files):
            if f.startswith(".") or f.endswith(SIDECAR_SUFFIX):
                continue
            yield os.path.join(root, f)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="*")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--delete", metavar="NAME")
    ap.add_argument("--force", action="store_true",
                    help="ingest even if it duplicates another document")
    args = ap.parse_args(argv)
    col = rag_utils.get_collection()

    if args.list:
        for d in list_documents(col):
            print(f"{d['chunks']:6}  {d['added_at']:19}  {d['source']}  ({d['file_path']})")
        return 0
    if args.delete:
        print(f"Removed {delete_document(col, args.delete)} document(s) named {args.delete}")
        return 0
    if not args.paths:
        ap.print_help()
        return 2

    failed = 0
    for top in args.paths:
        for path in _files_under(top):
            r = ingest_file(path, col, force=args.force)
            failed += r.status == "error"
            note = f"{r.chunks} chunks" if r.status == "ingested" else r.detail
            print(f"{r.status:18}  {r.source}  {note}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
