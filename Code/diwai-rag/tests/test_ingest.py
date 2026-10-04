import hashlib
import json
import os
import re
import subprocess

import pytest

import config
import ingest
import rag_utils
from conftest import fake_vector


@pytest.fixture
def col(temp_store):
    return rag_utils.get_collection()


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def long_text(tag, n=3000):
    return " ".join(f"{tag}{i}" for i in range(n))


def chunks_of(col, doc_id):
    return col.get(where={"doc_id": doc_id}, include=["metadatas"])


# ---- basic ingest and metadata --------------------------------------------

def test_ingest_text_file_records_metadata(col, temp_store):
    f = write(temp_store / "in" / "notes.txt", long_text("w"))
    r = ingest.ingest_file(str(f), col)
    assert r.status == "ingested"
    assert r.source == "notes.txt"
    assert r.chunks == col.count() > 1
    meta = col.get(include=["metadatas"])["metadatas"][0]
    assert meta["source"] == "notes.txt"
    assert meta["file_path"] == os.path.realpath(str(f))
    assert meta["doc_id"] == ingest.doc_id_for(str(f))
    assert re.fullmatch(r"[0-9a-f]{64}", meta["content_sha256"])
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d", meta["added_at"])
    ids = col.get()["ids"]
    assert f"{meta['doc_id']}::chunk::0" in ids


def test_doc_id_is_full_path_and_follows_symlinks(temp_store):
    a = write(temp_store / "a" / "x.txt", "one")
    b = write(temp_store / "b" / "x.txt", "two")
    link = temp_store / "link.txt"
    link.symlink_to(a)
    assert ingest.doc_id_for(str(a)) != ingest.doc_id_for(str(b))
    assert ingest.doc_id_for(str(link)) == ingest.doc_id_for(str(a))
    assert len(ingest.doc_id_for(str(a))) == 16


def test_progress_callback_reports_each_batch(col, temp_store, monkeypatch):
    monkeypatch.setattr(config, "EMBED_BATCH", 2)
    f = write(temp_store / "p.txt", long_text("p"))
    calls = []
    r = ingest.ingest_file(str(f), col, progress=lambda d, t: calls.append((d, t)))
    assert len(calls) == -(-r.chunks // 2)
    assert calls[-1] == (r.chunks, r.chunks)


def test_small_upsert_batches_store_every_chunk(col, temp_store, monkeypatch):
    monkeypatch.setattr(config, "UPSERT_BATCH", 2)
    f = write(temp_store / "u.txt", long_text("u"))
    r = ingest.ingest_file(str(f), col)
    assert col.count() == r.chunks > 2


# ---- identity, replace, dedupe --------------------------------------------

def test_same_name_in_two_folders_both_survive(col, temp_store):
    a = write(temp_store / "a" / "Report.txt", long_text("alpha"))
    b = write(temp_store / "b" / "Report.txt", long_text("beta"))
    ra = ingest.ingest_file(str(a), col)
    rb = ingest.ingest_file(str(b), col)
    assert (ra.status, rb.status) == ("ingested", "ingested")
    assert col.count() == ra.chunks + rb.chunks


def test_reingest_replaces_exactly_with_no_tails(col, temp_store):
    f = write(temp_store / "r.txt", long_text("old", 4000))
    first = ingest.ingest_file(str(f), col)
    write(f, long_text("new", 1000))
    second = ingest.ingest_file(str(f), col)
    assert second.status == "ingested"
    assert second.chunks < first.chunks
    got = chunks_of(col, ingest.doc_id_for(str(f)))
    assert len(got["ids"]) == second.chunks == col.count()


def test_exact_duplicate_elsewhere_is_skipped(col, temp_store):
    a = write(temp_store / "a" / "orig.txt", long_text("d"))
    b = write(temp_store / "b" / "copy.txt", long_text("d"))
    ingest.ingest_file(str(a), col)
    before = col.count()
    r = ingest.ingest_file(str(b), col)
    assert r.status == "skipped_duplicate"
    assert r.duplicate_of == "orig.txt"
    assert col.count() == before


def test_whitespace_only_difference_is_a_duplicate(col, temp_store):
    a = write(temp_store / "a.txt", "alpha  beta\ngamma")
    b = write(temp_store / "b.txt", "alpha beta   gamma\n\n")
    ingest.ingest_file(str(a), col)
    assert ingest.ingest_file(str(b), col).status == "skipped_duplicate"


def test_force_ingests_a_duplicate(col, temp_store):
    a = write(temp_store / "a.txt", long_text("f"))
    b = write(temp_store / "b.txt", long_text("f"))
    ingest.ingest_file(str(a), col)
    assert ingest.ingest_file(str(b), col, force=True).status == "ingested"


def _digit_blind_embedder(monkeypatch):
    monkeypatch.setattr(
        rag_utils, "embed_texts",
        lambda texts: [fake_vector(re.sub(r"\d", "", t)) for t in texts],
    )


def test_near_duplicate_is_skipped(col, temp_store, monkeypatch):
    _digit_blind_embedder(monkeypatch)
    base = long_text("word", 3000)
    a = write(temp_store / "v1.txt", "release 1 " + base)
    b = write(temp_store / "v2.txt", "release 2 " + base)
    ingest.ingest_file(str(a), col)
    r = ingest.ingest_file(str(b), col)
    assert r.status == "skipped_duplicate"
    assert r.duplicate_of == "v1.txt"


def test_partial_overlap_is_not_a_near_duplicate(col, temp_store, monkeypatch):
    _digit_blind_embedder(monkeypatch)
    shared = long_text("shared", 400)
    a = write(temp_store / "a.txt", shared + " " + long_text("aaa", 3000))
    b = write(temp_store / "b.txt", shared + " " + long_text("bbb", 3000))
    ingest.ingest_file(str(a), col)
    assert ingest.ingest_file(str(b), col).status == "ingested"


def test_embedding_failure_keeps_the_old_copy(col, temp_store, monkeypatch):
    f = write(temp_store / "k.txt", long_text("keep"))
    first = ingest.ingest_file(str(f), col)
    write(f, long_text("changed"))

    def boom(texts):
        raise ConnectionError("LM Studio down")

    monkeypatch.setattr(rag_utils, "embed_texts", boom)
    r = ingest.ingest_file(str(f), col)
    assert r.status == "error"
    assert "LM Studio down" in r.detail
    assert col.count() == first.chunks


# ---- empty / unsupported --------------------------------------------------

def test_empty_file_is_reported_empty(col, temp_store):
    f = write(temp_store / "e.txt", "   \n")
    r = ingest.ingest_file(str(f), col)
    assert r.status == "empty"
    assert col.count() == 0


def test_unsupported_type_is_reported(col, temp_store):
    f = write(temp_store / "x.xyz", "data")
    r = ingest.ingest_file(str(f), col)
    assert r.status == "error"
    assert "Unsupported" in r.detail


def test_upper_case_extension_is_read(temp_store):
    f = write(temp_store / "README.TXT", "hello")
    assert ingest.read_file(str(f)) == "hello"


# ---- readers --------------------------------------------------------------

def _pdf_with_text(path, text):
    ps = f"/Helvetica findfont 12 scalefont setfont 72 720 moveto ({text}) show showpage"
    subprocess.run(["gs", "-q", "-sDEVICE=pdfwrite", "-o", str(path), "-c", ps],
                   check=True)
    return path


def test_pdf_text_is_read(temp_store):
    f = _pdf_with_text(temp_store / "t.pdf", "Hello chrony makestep")
    assert "Hello chrony makestep" in ingest.read_file(str(f))


def test_pdf_falls_back_to_ghostscript_when_pypdf_raises(temp_store, monkeypatch):
    f = _pdf_with_text(temp_store / "t.pdf", "Fallback works")

    def broken(path):
        raise ValueError("pypdf cannot parse")

    monkeypatch.setattr(ingest, "_pdf_text_pypdf", broken)
    assert "Fallback works" in ingest.read_file(str(f))


def test_image_only_pdf_needs_ocr(col, temp_store):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    f = temp_store / "scan.pdf"
    with open(f, "wb") as fh:
        w.write(fh)
    r = ingest.ingest_file(str(f), col)
    assert r.status == "empty"
    assert r.detail == "No text found (needs OCR)"


def test_docx_reads_tables_in_document_order(temp_store):
    import docx
    d = docx.Document()
    d.add_paragraph("Before table")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "CellOne"
    t.cell(0, 1).text = "CellTwo"
    d.add_paragraph("After table")
    f = temp_store / "d.docx"
    d.save(f)
    text = ingest.read_file(str(f))
    assert text.index("Before table") < text.index("CellOne") < text.index("After table")
    assert "CellTwo" in text


def test_pptx_reads_slide_text_and_notes(temp_store):
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Slide Title"
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Body text"
    s.notes_slide.notes_text_frame.text = "Speaker notes"
    f = temp_store / "s.pptx"
    p.save(f)
    text = ingest.read_file(str(f))
    for want in ("Slide Title", "Body text", "Speaker notes"):
        assert want in text


def test_html_drops_scripts_and_styles(temp_store):
    f = write(temp_store / "p.html",
              "<html><head><style>.x{}</style><script>evil()</script></head>"
              "<body><h1>Heading</h1><p>Para</p></body></html>")
    text = ingest.read_file(str(f))
    assert "Heading" in text and "Para" in text
    assert "evil" not in text and ".x{}" not in text


def test_dita_keeps_title(temp_store):
    f = write(temp_store / "c.dita",
              '<?xml version="1.0"?><concept id="c"><title>3.13.11 FIPS crypto</title>'
              "<conbody><p>Use validated modules.</p></conbody></concept>")
    text = ingest.read_file(str(f))
    assert "3.13.11 FIPS crypto" in text and "Use validated modules." in text


def test_xml_external_entities_are_not_resolved(temp_store):
    secret = write(temp_store / "secret.txt", "TOPSECRET")
    f = write(temp_store / "x.dita",
              f'<?xml version="1.0"?><!DOCTYPE c [<!ENTITY e SYSTEM "file://{secret}">]>'
              "<concept><title>T</title><p>&e;</p></concept>")
    assert "TOPSECRET" not in ingest.read_file(str(f))


# ---- provenance sidecar ---------------------------------------------------

def test_provenance_sidecar_is_stored_and_verified(col, temp_store):
    f = write(temp_store / "guide.md", "chrony makestep guide")
    sha = hashlib.sha256(f.read_bytes()).hexdigest()
    write(temp_store / "guide.md.source.json", json.dumps(
        {"url": "https://example.org/chrony", "fetched": "2026-10-03", "sha256": sha}))
    ingest.ingest_file(str(f), col)
    meta = col.get(include=["metadatas"])["metadatas"][0]
    assert meta["provenance_url"] == "https://example.org/chrony"
    assert meta["provenance_fetched"] == "2026-10-03"
    assert meta["provenance_verified"] is True


def test_provenance_hash_mismatch_is_flagged(col, temp_store):
    f = write(temp_store / "guide.md", "edited after fetch")
    write(temp_store / "guide.md.source.json", json.dumps(
        {"url": "u", "fetched": "2026-10-03", "sha256": "0" * 64}))
    ingest.ingest_file(str(f), col)
    meta = col.get(include=["metadatas"])["metadatas"][0]
    assert meta["provenance_verified"] is False


def test_no_sidecar_means_no_provenance_fields(col, temp_store):
    f = write(temp_store / "plain.md", "text")
    ingest.ingest_file(str(f), col)
    meta = col.get(include=["metadatas"])["metadatas"][0]
    assert "provenance_url" not in meta


# ---- list / delete / CLI --------------------------------------------------

def test_list_documents_groups_chunks_and_pages(col, temp_store, monkeypatch):
    monkeypatch.setattr(config, "UPSERT_BATCH", 3)  # also the page size
    a = write(temp_store / "a.txt", long_text("aa"))
    b = write(temp_store / "b.txt", "short")
    ra = ingest.ingest_file(str(a), col)
    ingest.ingest_file(str(b), col)
    docs = {d["source"]: d for d in ingest.list_documents(col)}
    assert docs["a.txt"]["chunks"] == ra.chunks
    assert docs["b.txt"]["chunks"] == 1
    assert docs["a.txt"]["file_path"] == os.path.realpath(str(a))


def test_delete_by_name_removes_every_chunk(col, temp_store):
    a = write(temp_store / "a.txt", long_text("aa"))
    b = write(temp_store / "b.txt", "keep me")
    ingest.ingest_file(str(a), col)
    ingest.ingest_file(str(b), col)
    assert ingest.delete_document(col, "a.txt") == 1
    assert [d["source"] for d in ingest.list_documents(col)] == ["b.txt"]
    assert ingest.delete_document(col, "a.txt") == 0


def test_cli_ingests_folder_skipping_sidecars_and_hidden(temp_store, capsys):
    folder = temp_store / "drop"
    write(folder / "one.txt", "first doc")
    write(folder / "sub" / "two.md", "second doc")
    write(folder / "one.txt.source.json", "{}")
    write(folder / ".DS_Store", "junk")
    assert ingest.main([str(folder)]) == 0
    out = capsys.readouterr().out
    assert "one.txt" in out and "two.md" in out
    assert "source.json" not in out and "DS_Store" not in out
    ingest.main(["--list"])
    listed = capsys.readouterr().out
    assert "one.txt" in listed and "two.md" in listed
    ingest.main(["--delete", "one.txt"])
    assert "Removed 1" in capsys.readouterr().out


def test_ingest_embeds_with_the_document_label(col, temp_store, embed_calls):
    f = write(temp_store / "l.txt", long_text("lab"))
    ingest.ingest_file(str(f), col)
    assert embed_calls
    assert all(t.startswith("search_document: ") for t in embed_calls)


def test_stored_chunk_text_has_no_label(col, temp_store):
    f = write(temp_store / "l.txt", "plain words")
    ingest.ingest_file(str(f), col)
    assert col.get()["documents"] == ["plain words"]


# ---- vendor web pages: keep the content, drop the site chrome --------------

PAGE = """<html><head><title>T</title></head><body>
<header>Site header Skip to navigation</header>
<nav><ul><li>Chapter list</li><li>Other chapter</li></ul></nav>
<main><h1>Troubleshooting chrony</h1><p>Run chronyc tracking.</p><p>Then chronyc sources.</p></main>
<aside>Was this helpful?</aside><footer>Copyright footer</footer></body></html>"""


def test_html_uses_main_content_and_drops_site_chrome(temp_store):
    f = write(temp_store / "page.html", PAGE)
    text = ingest.read_file(str(f))
    assert "Troubleshooting chrony" in text and "chronyc sources" in text
    for junk in ("Site header", "Chapter list", "Was this helpful", "Copyright footer"):
        assert junk not in text, junk


def test_html_role_main_counts_as_main(temp_store):
    f = write(temp_store / "s.html", '<html><body><div role="navigation">Menu</div>'
              '<div role="main"><p>Rule reload</p></div></body></html>')
    text = ingest.read_file(str(f))
    assert "Rule reload" in text and "Menu" not in text


def test_html_without_main_keeps_body_minus_chrome(temp_store):
    f = write(temp_store / "b.html", "<html><body><nav>Menu</nav><p>Body text</p>"
              "<footer>Foot</footer></body></html>")
    text = ingest.read_file(str(f))
    assert "Body text" in text and "Menu" not in text and "Foot" not in text


def test_html_keeps_paragraph_breaks(temp_store):
    f = write(temp_store / "p.html", "<html><body><main><p>First para.</p><p>Second para.</p>"
              "<pre>line one\nline two</pre></main></body></html>")
    text = ingest.read_file(str(f))
    assert "First para.\n\nSecond para." in text
    assert "line one\nline two" in text
